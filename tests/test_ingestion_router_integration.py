from __future__ import annotations

import importlib.util
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from src.artifacts import write_ingested_artifacts
from src.ingestion.parsing.lift import LiftAPIParserClient
from src.ingestion.parsing.service import ParsingService
from src.ingestion.runner import run as run_ingestion
from src.models import ParsedData, PipelineState


class IngestionRouterIntegrationTests(unittest.TestCase):
    def test_lift_provider_routes_markdown_to_local_text_parser(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            path = root / "README.md"
            path.write_text("# Routed locally\n\nContent.", encoding="utf-8")

            with patch.object(
                LiftAPIParserClient,
                "parse_file",
                side_effect=AssertionError("Lift must not receive Markdown"),
            ):
                output = run_ingestion(
                    path,
                    parser_config={"provider": "lift_api"},
                    project_root=root,
                )

        self.assertEqual(len(output.parsed_data), 1)
        parsed = output.parsed_data[0]
        self.assertEqual(parsed.metadata["parser"], "text")
        self.assertEqual(parsed.metadata["backend"], "text")
        self.assertEqual(parsed.text, "# Routed locally\n\nContent.")

    def test_chandra_provider_routes_csv_to_local_text_parser(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            path = root / "records.csv"
            path.write_text("name,score\nAda,10\n", encoding="utf-8")

            output = run_ingestion(
                path,
                parser_config={"provider": "chandra2"},
                project_root=root,
            )

        parsed = output.parsed_data[0]
        self.assertEqual(parsed.metadata["parser"], "text")
        self.assertEqual(parsed.metadata["backend"], "text")
        self.assertEqual(parsed.tables[0].headers, ["name", "score"])
        self.assertEqual(parsed.tables[0].rows, [["Ada", "10"]])

    @unittest.skipUnless(
        importlib.util.find_spec("openpyxl"),
        "openpyxl is not installed",
    )
    def test_lift_provider_routes_xlsx_to_local_table_parser(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            path = root / "records.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "People"
            worksheet.append(["name", "score"])
            worksheet.append(["Ada", 10])
            workbook.save(path)

            with patch.object(
                LiftAPIParserClient,
                "parse_file",
                side_effect=AssertionError("Lift must not receive XLSX"),
            ):
                output = run_ingestion(
                    path,
                    parser_config={"provider": "lift_api"},
                    project_root=root,
                )

        parsed = output.parsed_data[0]
        self.assertEqual(parsed.metadata["parser"], "table")
        self.assertEqual(parsed.metadata["backend"], "table")
        self.assertEqual(parsed.tables[0].name, "People")
        self.assertEqual(parsed.tables[0].rows, [["Ada", "10"]])

    def test_top_level_lift_provider_still_handles_visual_documents(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            path = root / "sample.pdf"
            path.write_bytes(b"%PDF-1.4\n")

            def fake_parse(
                _client: LiftAPIParserClient,
                source: str | Path,
                data_object: object,
            ) -> ParsedData:
                return ParsedData(
                    object_id=data_object.object_id,
                    source_uri=data_object.uri,
                    source_format="pdf",
                    rows=[{"extraction": {"main_text": "Lift output"}}],
                    text="Lift output",
                    metadata={"parser": "lift-api"},
                )

            with patch.object(
                LiftAPIParserClient,
                "parse_file",
                autospec=True,
                side_effect=fake_parse,
            ) as parse_file:
                output = run_ingestion(
                    path,
                    parser_config={"provider": "lift_api"},
                    project_root=root,
                )

        parse_file.assert_called_once()
        self.assertEqual(output.parsed_data[0].text, "Lift output")
        self.assertEqual(output.parsed_data[0].metadata["backend"], "lift_api")

    def test_ingested_document_contract_is_unchanged_after_routing(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            path = root / "note.txt"
            path.write_text("Stable ingestion output.", encoding="utf-8")
            output = run_ingestion(
                path,
                parser_config={"provider": "lift_api"},
                project_root=root,
            )
            state = PipelineState(
                run_id="router-test",
                input_source=path.as_posix(),
                embedded_dir="embedded",
                output_dir="output",
                data_objects=output.data_objects,
                parsed_data=output.parsed_data,
                initial_schemas=output.initial_schemas,
                ingestion_config={"provider": "lift_api"},
            )
            artifact_dir = root / "ingested"
            write_ingested_artifacts(state, artifact_dir, project_root=root)
            document_path = next((artifact_dir / "documents").glob("*.json"))
            payload = json.loads(document_path.read_text(encoding="utf-8"))

        self.assertEqual(payload["contract_version"], "ingested-document-v2")
        self.assertEqual(payload["stage"], "ingested")
        self.assertEqual(payload["status"], "succeeded")
        self.assertEqual(
            set(payload),
            {
                "contract_version",
                "stage",
                "status",
                "document_id",
                "source",
                "parsed",
                "schema_id",
                "failure",
            },
        )
        self.assertEqual(
            set(payload["parsed"]),
            {
                "object_id",
                "source_uri",
                "source_format",
                "rows",
                "text",
                "metadata",
                "tables",
            },
        )


class DocumentProviderConfigTests(unittest.TestCase):
    def test_top_level_chandra_provider_only_changes_document_route(self) -> None:
        router = ParsingService.from_config({"provider": "chandra2"}).router

        self.assertEqual(router.resolve(Path("scan.pdf")).provider_name, "chandra2")
        self.assertEqual(router.resolve(Path("note.md")).backend_name, "text")
        self.assertEqual(router.resolve(Path("book.xlsx")).backend_name, "table")

    def test_legacy_router_config_remains_supported(self) -> None:
        router = ParsingService.from_config(
            {
                "provider": "router",
                "document": {"provider": "chandra2"},
            }
        ).router

        self.assertEqual(router.resolve(Path("scan.pdf")).provider_name, "chandra2")


if __name__ == "__main__":
    unittest.main()
